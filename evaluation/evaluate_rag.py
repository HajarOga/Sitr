"""Evaluate Sitr's policy retriever without invoking the answer model.

Run from the project root:
    python -m evaluation.evaluate_rag

Each run uses a fresh in-memory Chroma collection and writes a timestamped workbook.
"""
from __future__ import annotations

import statistics
import sys
import uuid
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app import config  # noqa: E402
from app.policy import qa  # noqa: E402
from app.policy.policy_index import (  # noqa: E402
    PolicyIndex,
    SentenceTransformerEmbedder,
)
from app.rules.store import library as default_library  # noqa: E402

RESULTS_DIR = ROOT / "evaluation" / "results"
TOP_K_VALUES = (1, 3, 5, 8)

# Expected rule IDs were assigned from the checked-in policy. Receiver/type filters
# remain implicit except where noted, so the normal question heuristics are measured too.
CASES: tuple[dict[str, Any], ...] = (
    {"id": "en-vendor-name", "language": "English",
     "question": "May an external vendor see the customer's full name?",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "PERSON_NAME",
     "gold": ["VENDOR-PII-001"]},
    {"id": "en-vendor-national-id", "language": "English",
     "question": "Does the vendor need the customer's National ID?",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "NATIONAL_ID",
     "gold": ["VENDOR-PII-002"]},
    {"id": "en-vendor-phone", "language": "English",
     "question": "May a vendor call the customer using their phone number?",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "PHONE",
     "gold": ["VENDOR-PII-003"]},
    {"id": "en-vendor-email", "language": "English",
     "question": "Can an outside vendor retain the customer's email address?",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "EMAIL",
     "gold": ["VENDOR-PII-004"]},
    {"id": "en-vendor-address", "language": "English",
     "question": "Can the vendor see the customer's full home address?",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "ADDRESS",
     "gold": ["VENDOR-PII-005"]},
    {"id": "en-vendor-customer-id", "language": "English",
     "question": "Does the vendor need the customer ID to match a repair job?",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "CUSTOMER_ID",
     "gold": ["VENDOR-PII-010"]},
    {"id": "en-hr-passport", "language": "English",
     "question": "May HR keep passport details in an employee file?",
     "receiver": "INTERNAL_HR", "pii_type": "PASSPORT_NUMBER",
     "gold": ["HR-PII-009"]},
    {"id": "en-hr-employee-id", "language": "English",
     "question": "Does HR need the employee ID to process a record?",
     "receiver": "INTERNAL_HR", "pii_type": "EMPLOYEE_ID",
     "gold": ["HR-PII-011"]},
    {"id": "en-analyst-bank", "language": "English",
     "question": "Can an analyst use a customer's bank account number?",
     "receiver": "ANALYST", "pii_type": "BANK_ACCOUNT",
     "gold": ["ANALYST-PII-007"]},
    {"id": "en-analyst-customer-id", "language": "English",
     "question": "May analysts keep a pseudonymous customer ID to join tables?",
     "receiver": "ANALYST", "pii_type": "CUSTOMER_ID",
     "gold": ["ANALYST-PII-010"]},
    {"id": "en-support-phone", "language": "English",
     "question": "Can customer support retain a callback phone number?",
     "receiver": "CUSTOMER_SUPPORT", "pii_type": "PHONE",
     "gold": ["SUPPORT-PII-003"]},
    {"id": "en-support-national-id", "language": "English",
     "question": "Does customer support need a customer's national ID?",
     "receiver": "CUSTOMER_SUPPORT", "pii_type": "NATIONAL_ID",
     "gold": ["SUPPORT-PII-002"]},
    {"id": "ar-vendor-name", "language": "Arabic",
     "question": "هل يجوز للمورد الخارجي الاطلاع على اسم العميل؟",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "PERSON_NAME",
     "gold": ["VENDOR-PII-001"]},
    {"id": "ar-vendor-national-id", "language": "Arabic",
     "question": "هل يجوز للمورد الاطلاع على رقم الهوية الوطنية؟",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "NATIONAL_ID",
     "gold": ["VENDOR-PII-002"]},
    {"id": "ar-vendor-phone", "language": "Arabic",
     "question": "هل يمكن مشاركة رقم الهاتف مع المورد؟",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "PHONE",
     "gold": ["VENDOR-PII-003"]},
    {"id": "ar-vendor-email", "language": "Arabic",
     "question": "هل بريد العميل الإلكتروني ضروري للمورد الخارجي؟",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "EMAIL",
     "gold": ["VENDOR-PII-004"]},
    {"id": "ar-vendor-address", "language": "Arabic",
     "question": "هل يحتاج المورد إلى عنوان المنزل كاملًا؟",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "ADDRESS",
     "gold": ["VENDOR-PII-005"]},
    {"id": "ar-vendor-iban", "language": "Arabic",
     "question": "هل نشارك الآيبان مع المورد؟",
     "receiver": "EXTERNAL_VENDOR", "pii_type": "IBAN",
     "gold": ["VENDOR-PII-006"]},
    {"id": "ar-hr-name", "language": "Arabic",
     "question": "هل يجوز للموارد البشرية الاحتفاظ باسم الموظف؟",
     "receiver": "INTERNAL_HR", "pii_type": "PERSON_NAME",
     "gold": ["HR-PII-001"]},
    {"id": "ar-hr-employee-id", "language": "Arabic",
     "question": "هل يحتاج موظف الموارد البشرية إلى الرقم الوظيفي؟",
     "receiver": "INTERNAL_HR", "pii_type": "EMPLOYEE_ID",
     "gold": ["HR-PII-011"]},
    {"id": "ar-analyst-address", "language": "Arabic",
     "question": "هل يكفي اسم المدينة بدل العنوان الكامل لتحليل البيانات؟",
     "receiver": "ANALYST", "pii_type": "ADDRESS",
     "gold": ["ANALYST-PII-005"]},
    {"id": "ar-support-national-id", "language": "Arabic",
     "question": "هل يمكن للدعم طلب الهوية الوطنية من العميل؟",
     "receiver": "CUSTOMER_SUPPORT", "pii_type": "NATIONAL_ID",
     "gold": ["SUPPORT-PII-002"]},
    {"id": "ar-support-mobile", "language": "Arabic",
     "question": "هل يستطيع دعم العملاء الاتصال بالعميل عبر رقم الجوال؟",
     "receiver": "CUSTOMER_SUPPORT", "pii_type": "PHONE",
     "gold": ["SUPPORT-PII-003"]},
    {"id": "mixed-analyst-vat", "language": "Mixed",
     "question": "For ANALYST, can we keep the CR/VAT identifier in reports?",
     "receiver": "ANALYST", "pii_type": "COMPANY_IDENTIFIER",
     "gold": ["ANALYST-PII-012"]},
    {"id": "en-any-receiver-phone", "language": "English",
     "question": "Can a phone number be retained?",
     "receiver": None, "pii_type": "PHONE",
     "gold": ["VENDOR-PII-003", "HR-PII-003",
              "ANALYST-PII-003", "SUPPORT-PII-003"]},
    {"id": "en-custom-no-rules", "language": "English",
     "question": "May a custom receiver retain a name?",
     "receiver": "CUSTOM", "pii_type": "PERSON_NAME", "gold": []},
    {"id": "ar-custom-filtered-no-rules", "language": "Arabic",
     "question": "هل يمكن للمستقبل المخصص مشاركة الاسم؟",
     "receiver": "CUSTOM", "pii_type": "PERSON_NAME", "gold": [],
     "receiver_override": "CUSTOM"},
    {"id": "en-custom-out-of-domain", "language": "English",
     "question": "Does the custom receiver have a rule about retina scans?",
     "receiver": "CUSTOM", "pii_type": None, "gold": []},
    {"id": "en-out-of-domain", "language": "English",
     "question": "What is the weather in Riyadh?",
     "receiver": None, "pii_type": None, "gold": []},
    {"id": "edge-empty", "language": "Edge",
     "question": "   ", "receiver": None, "pii_type": None, "gold": []},
)


def _ratio(numerator: float, denominator: float) -> float | None:
    return numerator / denominator if denominator else None


def _candidate_rules(lib, receiver: str | None, pii_type: str | None) -> int:
    return sum(
        (receiver is None or rule.receiver == receiver)
        and (pii_type is None or rule.pii_type == pii_type)
        for rule in lib.rules
    )


def _run_case(case: dict[str, Any], top_k: int, index: PolicyIndex, lib) -> dict[str, Any]:
    kwargs: dict[str, Any] = {}
    if case.get("receiver_override") is not None:
        kwargs["receiver"] = case["receiver_override"]
    try:
        result = qa.ask_policy(
            case["question"], top_k=top_k, use_model=False, index=index,
            library=lib, **kwargs,
        )
    except Exception as exc:  # Store per-case errors and continue the benchmark.
        return {
            "case_id": case["id"], "language": case["language"], "top_k": top_k,
            "question": case["question"], "expected_receiver": case["receiver"],
            "detected_receiver": "", "expected_type": case["pii_type"],
            "detected_type": "", "expected_rule_ids": ", ".join(case["gold"]),
            "retrieved_rule_ids": "", "ranks": "", "distances": "",
            "relevance_scores": "", "actions": "", "citations": "",
            "grounded_answer": "", "notes": "",
            "answerable": bool(case["gold"]), "retrieved_relevant": 0,
            "retrieved_count": 0, "precision_at_k": None, "recall_at_k": None,
            "hit_rate": None, "reciprocal_rank": None, "faithful_evidence_rate": None,
            "citation_accuracy": None, "citation_coverage": None,
            "unsupported_hits": 0, "false_positive": None,
            "receiver_filter_correct": None, "type_filter_correct": None,
            "canonical_source_consistency": None,
            "error": f"{type(exc).__name__}: {exc}",
        }

    hits = result.get("hits", [])
    ids = [hit["rule_id"] for hit in hits]
    gold = set(case["gold"])
    retrieved_relevant = sum(rule_id in gold for rule_id in ids)
    reciprocal_rank = next(
        (1 / rank for rank, rule_id in enumerate(ids, start=1) if rule_id in gold),
        0.0 if gold else None,
    )
    actual_receiver = result.get("receiver_filter")
    actual_type = result.get("type_filter")
    denominator = max(1, min(top_k, _candidate_rules(lib, actual_receiver, actual_type)))
    precision = retrieved_relevant / denominator if gold else None
    recall = _ratio(retrieved_relevant, len(gold)) if gold else None

    rule_by_id = {rule.rule_id: rule for rule in lib.rules}
    canonical_hits = 0
    for hit in hits:
        source = rule_by_id.get(hit["rule_id"])
        if source and all(hit.get(field) == getattr(source, field)
                          for field in ("receiver", "pii_type", "action", "reason", "citation")):
            canonical_hits += 1
    faithfulness = (
        retrieved_relevant / len(hits) if hits
        else (1.0 if not gold else 0.0)
    )

    citations = result.get("citations", [])
    valid_citations = sum(
        citation.get("rule_id") in ids
        and citation.get("rule_id") in rule_by_id
        and citation.get("citation") == rule_by_id[citation["rule_id"]].citation
        and citation.get("action") == rule_by_id[citation["rule_id"]].action
        for citation in citations
    )
    citation_accuracy = _ratio(valid_citations, len(citations))
    citation_coverage = _ratio(
        len(gold.intersection(citation.get("rule_id") for citation in citations)), len(gold)
    ) if gold else None
    unsupported_hits = sum(rule_id not in gold for rule_id in ids)
    false_positive = bool(hits) if not gold else None
    receiver_correct = actual_receiver == case["receiver"]
    type_correct = actual_type == case["pii_type"]
    answer = qa.grounded_answer(result, hits, lib)
    return {
        "case_id": case["id"], "language": case["language"], "top_k": top_k,
        "question": case["question"], "expected_receiver": case["receiver"],
        "detected_receiver": actual_receiver, "expected_type": case["pii_type"],
        "detected_type": actual_type, "expected_rule_ids": ", ".join(case["gold"]),
        "retrieved_rule_ids": ", ".join(ids),
        "ranks": ", ".join(f"{i + 1}:{rule_id}" for i, rule_id in enumerate(ids)),
        "distances": ", ".join(str(hit.get("distance")) for hit in hits),
        "relevance_scores": ", ".join(str(hit.get("relevance")) for hit in hits),
        "actions": "; ".join(f"{hit['rule_id']}={hit['action']}" for hit in hits),
        "citations": "; ".join(
            f"{item.get('rule_id')} | {item.get('citation')}" for item in citations
        ),
        "grounded_answer": answer, "notes": " | ".join(result.get("notes", [])),
        "answerable": bool(gold), "retrieved_relevant": retrieved_relevant,
        "retrieved_count": len(hits), "precision_at_k": precision,
        "recall_at_k": recall,
        "hit_rate": int(retrieved_relevant > 0) if gold else None,
        "reciprocal_rank": reciprocal_rank,
        "faithful_evidence_rate": faithfulness,
        "citation_accuracy": citation_accuracy, "citation_coverage": citation_coverage,
        "unsupported_hits": unsupported_hits, "false_positive": false_positive,
        "receiver_filter_correct": receiver_correct, "type_filter_correct": type_correct,
        "canonical_source_consistency": (
            canonical_hits / len(hits) if hits else (1.0 if not gold else 0.0)
        ),
        "error": result.get("error", ""),
    }


def _aggregate(rows: list[dict[str, Any]], field: str) -> float | None:
    values = [row[field] for row in rows if row.get(field) is not None]
    return statistics.mean(values) if values else None


def _style_sheet(sheet, max_width: int = 48) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    for column in sheet.columns:
        index = column[0].column
        longest = max((len(str(cell.value or "")) for cell in column[:300]), default=10)
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(longest + 2, 12), max_width
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def evaluate() -> Path:
    lib = default_library()
    known_ids = {rule.rule_id for rule in lib.rules}
    unknown_gold = sorted({
        rule_id for case in CASES for rule_id in case["gold"] if rule_id not in known_ids
    })
    if unknown_gold:
        raise ValueError(f"Benchmark contains rule IDs absent from the policy: {unknown_gold}")

    embedder = SentenceTransformerEmbedder()
    if not embedder.is_installed():
        raise RuntimeError(
            f"Configured embedding model is not installed: {embedder.name()}. "
            "Install it using the project's model setup instructions before evaluating."
        )

    index = PolicyIndex(
        embedder=embedder, persist=False,
        collection_name=f"rag-eval-{uuid.uuid4().hex[:12]}",
        library=lib,
    )
    rows = [
        _run_case(case, top_k, index, lib)
        for case in CASES for top_k in TOP_K_VALUES
    ]

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output_path = RESULTS_DIR / f"safeshare_rag_{timestamp}.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Metrics"
    summary_headers = [
        "Language", "Top K", "Cases", "Answerable cases", "P@K",
        "Recall@K", "Hit Rate", "MRR", "Evidence faithfulness",
        "Citation accuracy", "Citation coverage", "Canonical source consistency",
        "Unsupported hit rate", "No-rule false-positive rate",
        "Receiver filter accuracy", "Type filter accuracy", "Errors",
    ]
    summary.append(summary_headers)
    languages = ["All", "English", "Arabic", "Mixed", "Edge"]
    for language in languages:
        for top_k in TOP_K_VALUES:
            selected = [row for row in rows if row["top_k"] == top_k and
                        (language == "All" or row["language"] == language)]
            answerable = [row for row in selected if row["answerable"] and not row["error"]]
            unanswerable = [row for row in selected if not row["answerable"] and not row["error"]]
            hit_total = sum(row["retrieved_count"] for row in selected if not row["error"])
            unsupported_total = sum(row["unsupported_hits"] for row in selected if not row["error"])
            false_positives = sum(bool(row["false_positive"]) for row in unanswerable)
            summary.append([
                language, top_k, len(selected), len(answerable),
                _aggregate(answerable, "precision_at_k"),
                _aggregate(answerable, "recall_at_k"),
                _aggregate(answerable, "hit_rate"),
                _aggregate(answerable, "reciprocal_rank"),
                _aggregate([r for r in selected if not r["error"]], "faithful_evidence_rate"),
                _aggregate([r for r in selected if not r["error"]], "citation_accuracy"),
                _aggregate(answerable, "citation_coverage"),
                _aggregate([r for r in selected if not r["error"]],
                           "canonical_source_consistency"),
                _ratio(unsupported_total, hit_total),
                _ratio(false_positives, len(unanswerable)),
                _aggregate([r for r in selected if not r["error"]],
                           "receiver_filter_correct"),
                _aggregate([r for r in selected if not r["error"]],
                           "type_filter_correct"),
                sum(bool(row["error"]) for row in selected),
            ])

    detail = workbook.create_sheet("Case Details")
    detail_headers = list(rows[0].keys())
    detail.append(detail_headers)
    for row in rows:
        detail.append([row.get(header) for header in detail_headers])

    definitions = workbook.create_sheet("Method & Notes")
    info = [
        ("Setting", "Value"),
        ("Evaluation timestamp", timestamp),
        ("Policy version", lib.version),
        ("Policy rules", len(lib.rules)),
        ("Embedding model", embedder.name()),
        ("Index storage", "Fresh in-memory Chroma collection; persistent index not read or changed"),
        ("Answer model", "Not called; ask_policy(use_model=False) for every case"),
        ("Top-K values", ", ".join(map(str, TOP_K_VALUES))),
        ("Benchmark cases", len(CASES)),
        ("Languages", "English, Arabic, mixed-script English/acronyms, empty-input edge case"),
        ("Gold judgments", "Manually specified relevant policy rule IDs in evaluate_rag.py"),
        ("P@K definition", "Relevant hits / min(K, policy candidate count after detected filters); answerable cases only"),
        ("Recall@K definition", "Relevant retrieved gold IDs / all expected gold IDs; answerable cases only"),
        ("Hit Rate", "Fraction of answerable queries with at least one relevant hit"),
        ("MRR", "Mean reciprocal rank of first relevant hit; answerable cases only"),
        ("Evidence faithfulness", "Fraction of retrieved rules that are relevant to the question according to manually specified gold IDs; empty expected results are faithful only when retrieval is empty"),
        ("Canonical source consistency", "Fraction of retrieved rules whose receiver, type, action, reason, and citation exactly match the policy record for that ID"),
        ("Citation accuracy", "Citations that refer to retrieved canonical rules with exact action and source / emitted citations"),
        ("Citation coverage", "Expected gold IDs cited / expected gold IDs; answerable cases only"),
        ("Unsupported hit rate", "Retrieved hits not in a case's gold set / all retrieved hits; retrieval-level false evidence proxy"),
        ("No-rule false-positive rate", "No-gold cases that returned at least one rule / no-gold cases"),
        ("LLM hallucinations", "Not assessed: the answer model is deliberately excluded. Unsupported retrieval and citation fabrication proxies are reported instead."),
        ("Errors", "Per-query exceptions are retained in Case Details and counted in Metrics."),
        ("Recommendations", "See the Recommendations sheet for measured trade-offs and case-specific improvement opportunities."),
    ]
    for record in info:
        definitions.append(record)

    recommendations = workbook.create_sheet("Recommendations")
    recommendations.append(["Priority", "Measured observation", "Recommendation"])
    mismatch_cases = {}
    for row in rows:
        if row["error"] or row["language"] != "Arabic":
            continue
        if not row["receiver_filter_correct"] or not row["type_filter_correct"]:
            mismatch_cases[row["case_id"]] = row
    if mismatch_cases:
        for case_id, row in mismatch_cases.items():
            recommendations.append([
                "High",
                f"{case_id}: detected receiver/type "
                f"{row['detected_receiver'] or 'none'}/{row['detected_type'] or 'none'}; "
                f"expected {row['expected_receiver'] or 'none'}/{row['expected_type'] or 'none'}. "
                f"Retrieved {row['retrieved_rule_ids'] or 'no rules'} for gold "
                f"{row['expected_rule_ids'] or 'no rules'}.",
                "Improve Arabic type-intent disambiguation and add paraphrase/contrast tests "
                "for ambiguous words such as اسم (name) in address questions.",
            ])

    no_rule_cases = {case["id"]: case for case in CASES if not case["gold"]}
    false_positive_cases = {
        row["case_id"]: row for row in rows
        if row["top_k"] == max(TOP_K_VALUES) and row["false_positive"]
    }
    no_rule_rate = _ratio(len(false_positive_cases), len(no_rule_cases))
    if false_positive_cases:
        recommendations.append([
            "High",
            f"{len(false_positive_cases)}/{len(no_rule_cases)} no-rule cases returned results "
            f"at K={max(TOP_K_VALUES)} ({no_rule_rate:.1%}); cases: "
            + ", ".join(false_positive_cases),
            "Evaluate a calibrated relevance/abstention threshold or explicit no-policy-match "
            "path against both in-domain and out-of-domain questions before changing behavior.",
        ])

    k5 = next((row for row in summary.iter_rows(min_row=2, values_only=True)
               if row[0] == "All" and row[1] == 5), None)
    k8 = next((row for row in summary.iter_rows(min_row=2, values_only=True)
               if row[0] == "All" and row[1] == 8), None)
    if k5 and k8:
        recommendations.append([
            "Medium",
            f"Overall recall is {k5[5]:.1%} at K=5 and {k8[5]:.1%} at K=8; "
            f"unsupported-hit rate rises from {k5[12]:.1%} to {k8[12]:.1%}.",
            "Treat K=5 as the better measured trade-off in this set; validate on a larger "
            "policy-question set before changing the configured default.",
        ])
    recommendations.append([
        "Scope",
        "The answer model was not called; citation accuracy/source consistency measure "
        "retrieval and code-built citations only.",
        "Run a separate, controlled answer-generation evaluation for model-written claims "
        "and hallucinations; do not interpret this retrieval-only run as an LLM score.",
    ])

    errors = [row for row in rows if row["error"]]
    error_sheet = workbook.create_sheet("Errors")
    error_sheet.append(["Case ID", "Language", "Top K", "Question", "Error"])
    for row in errors:
        error_sheet.append(
            [row["case_id"], row["language"], row["top_k"], row["question"], row["error"]]
        )
    if not errors:
        error_sheet.append(["None", "", "", "", "No benchmark query errors."])

    for sheet in workbook.worksheets:
        _style_sheet(sheet, 60 if sheet.title == "Case Details" else 48)
    for row in summary.iter_rows(min_row=2, min_col=5, max_col=16):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "0.00%"
    workbook.save(output_path)
    return output_path


if __name__ == "__main__":
    report = evaluate()
    print(f"RAG evaluation report: {report}")
