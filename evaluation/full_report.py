"""Run the complete second-set evaluation and write timestamped JSONL and Excel reports.

Run from the project root:
    python -m evaluation.full_report
"""
from __future__ import annotations

import argparse
import json
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import config
from app.llm import engine, prompts
from app.rules.store import library as default_library
from evaluation import metrics
from evaluation.evaluate import test_cases
from evaluation.reference import action_expected, must_not_be_visible

ROOT = Path(__file__).resolve().parents[1]
TEST_DIR = ROOT / "data" / "test_second"
RESULTS_DIR = ROOT / "evaluation" / "results"
EXPECTED_DOCUMENTS = 20
EXPECTED_LABELS = 87


def _match_findings(findings: list[dict], labels: list[dict]) -> dict[int, int]:
    """One-to-one match valid findings to labels by PII type and covered value."""
    used: set[int] = set()
    matches: dict[int, int] = {}
    for label_index, label in enumerate(labels):
        candidates = []
        for finding_index, finding in enumerate(findings):
            if finding_index in used or not metrics.is_structured(finding):
                continue
            if str(finding.get("type", "")).upper() != str(label["type"]).upper():
                continue
            reported = str(finding.get("original_text") or "")
            if metrics.covers(reported, str(label["text"])):
                exact = metrics.same_value(reported, str(label["text"]))
                candidates.append((not exact, finding_index))
        if candidates:
            _, finding_index = min(candidates)
            matches[label_index] = finding_index
            used.add(finding_index)
    return matches


def _ratio(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def _style_sheet(sheet, widths: dict[int, int] | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    if widths:
        for column, width in widths.items():
            sheet.column_dimensions[get_column_letter(column)].width = width
    else:
        for column in sheet.columns:
            index = column[0].column
            longest = max((len(str(cell.value or "")) for cell in column[:1000]), default=10)
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 48)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_workbook(path: Path, report: dict[str, Any]) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append([f"SafeShare {report.get('dataset_name', 'test_second')} evaluation"])
    summary["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="17365D")
    summary.merge_cells("A1:E1")
    summary.append(["Generated", report["generated_at"]])
    summary.append(["Dataset", str(report.get("test_dir", TEST_DIR))])
    summary.append(["Raw predictions", str(report["json_path"])])
    summary.append([])
    summary.append(["Model configuration", "Value"])
    for key, value in report["model"].items():
        summary.append([key, value])
    summary.append([])
    summary.append(["Evaluation metrics", "Value", "Numerator", "Denominator", "Definition"])
    for metric in report["summary_metrics"]:
        summary.append([
            metric["name"], metric["value"], metric.get("numerator"),
            metric.get("denominator"), metric["definition"]])
        if metric.get("format") == "percent":
            summary.cell(summary.max_row, 2).number_format = "0.00%"
    summary.column_dimensions["A"].width = 35
    summary.column_dimensions["B"].width = 58
    summary.column_dimensions["C"].width = 16
    summary.column_dimensions["D"].width = 16
    summary.column_dimensions["E"].width = 90
    for row_number in (6, 6 + len(report["model"]) + 2):
        for cell in summary[row_number]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="17365D")
    summary.freeze_panes = "A7"
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    documents = workbook.create_sheet("Document Results")
    doc_headers = [
        "Status", "Document ID", "Receiver", "Labeled PII", "Predicted findings",
        "True positives", "False positives", "Missed PII", "Precision", "Recall",
        "F1", "Action agreement", "Leak rate", "Deferred rate", "Inference seconds",
        "Tokens in", "Tokens out", "Failure",
    ]
    documents.append(doc_headers)
    for row in report["document_rows"]:
        documents.append([row.get(header) for header in doc_headers])
    for row_number in range(2, documents.max_row + 1):
        for column in (9, 10, 11, 12, 13, 14):
            documents.cell(row_number, column).number_format = "0.00%"
    _style_sheet(documents, {
        1: 13, 2: 28, 3: 22, 4: 13, 5: 18, 6: 14, 7: 16, 8: 13,
        9: 13, 10: 13, 11: 13, 12: 18, 13: 13, 14: 15, 15: 18,
        16: 12, 17: 12, 18: 55,
    })

    pii_sheet = workbook.create_sheet("PII-Level Results")
    pii_headers = [
        "Document ID", "Document status", "Receiver", "PII value", "PII type",
        "Expected action", "Predicted value", "Predicted type", "Predicted action",
        "Finding correct", "Action correct", "Citation valid", "Prediction reason",
    ]
    pii_sheet.append(pii_headers)
    for row in report["pii_rows"]:
        pii_sheet.append([row.get(header) for header in pii_headers])
    _style_sheet(pii_sheet, {
        1: 28, 2: 17, 3: 22, 4: 34, 5: 20, 6: 18, 7: 36, 8: 20,
        9: 19, 10: 16, 11: 16, 12: 16, 13: 55,
    })

    errors = workbook.create_sheet("Errors")
    error_headers = [
        "Document ID", "Error category", "PII value", "PII type",
        "Expected action", "Predicted action", "Details",
    ]
    errors.append(error_headers)
    for row in report["error_rows"]:
        errors.append([row.get(header) for header in error_headers])
    _style_sheet(errors, {1: 28, 2: 22, 3: 34, 4: 20, 5: 18, 6: 19, 7: 70})

    calls = workbook.create_sheet("Model Calls")
    call_headers = [
        "Document ID", "Purpose", "Device", "Seconds", "Tokens in", "Tokens out",
        "Tokens/second", "Prompt chars", "Output chars", "JSON attempt",
        "JSON valid", "JSON validation seconds",
    ]
    calls.append(call_headers)
    for row in report["call_rows"]:
        calls.append([row.get(header) for header in call_headers])
    _style_sheet(calls, {1: 28, 2: 28, 3: 14, 4: 14, 5: 12, 6: 12, 7: 16,
                         8: 14, 9: 14, 10: 14, 11: 12, 12: 24})

    workbook.save(path)


def run(test_dir: Path = TEST_DIR, output_dir: Path = RESULTS_DIR, *,
        dataset_name: str | None = None,
        expected_documents: int | None = EXPECTED_DOCUMENTS,
        expected_labels: int | None = EXPECTED_LABELS) -> tuple[Path, Path]:
    if engine.provider() != "local":
        raise RuntimeError("Evaluation requires SAFESHARE2_PROVIDER=local.")
    if config.MODEL_NAME != "Qwen/Qwen2.5-3B-Instruct":
        raise RuntimeError(f"Unexpected base model: {config.MODEL_NAME}")
    if not config.ADAPTER or not Path(config.ADAPTER).is_dir():
        raise RuntimeError(f"Required LoRA adapter is missing: {config.ADAPTER!r}")

    import torch

    if not torch.cuda.is_available():
        raise RuntimeError("CUDA is unavailable; refusing to evaluate on CPU.")
    if engine.resolve_device() != "cuda":
        raise RuntimeError(f"Expected CUDA, configured device is {engine.resolve_device()!r}.")
    if not torch.cuda.is_bf16_supported():
        raise RuntimeError("This validated run requires BF16 support on the CUDA GPU.")
    config.MODEL_DTYPE = "bfloat16"
    test_dir = test_dir.resolve()
    output_dir = output_dir.resolve()
    dataset_name = dataset_name or test_dir.name
    dataset_slug = re.sub(r"[^A-Za-z0-9_-]+", "_", dataset_name).strip("_")
    if not dataset_slug:
        raise ValueError("dataset_name must contain letters or numbers.")
    cases = test_cases(test_dir)
    label_total = sum(len(case["items"]) for case in cases)
    if ((expected_documents is not None and len(cases) != expected_documents)
            or (expected_labels is not None and label_total != expected_labels)):
        raise RuntimeError(
            f"Expected {expected_documents} documents/{expected_labels} labels; "
            f"found {len(cases)}/{label_total}.")

    output_dir.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    json_path = output_dir / f"safeshare_{dataset_slug}_{stamp}_predictions.jsonl"
    xlsx_path = output_dir / f"safeshare_{dataset_slug}_{stamp}_report.xlsx"
    started = time.perf_counter()
    lib = default_library()
    document_rows: list[dict] = []
    pii_rows: list[dict] = []
    error_rows: list[dict] = []
    call_rows: list[dict] = []
    doc_scores = []

    with json_path.open("x", encoding="utf-8") as raw_file:
        print(f"Preflight: CUDA {torch.cuda.get_device_name(0)}; "
              f"{len(cases)} documents / {label_total} labels.", flush=True)
        print(f"Loading {config.ACTIVE_MODEL} with adapter {config.ADAPTER}...", flush=True)
        load_started = time.perf_counter()
        startup_error = None
        try:
            engine.ensure_installed()
            engine.load()
        except Exception as exc:  # record startup failures against each document below
            startup_error = f"{type(exc).__name__}: {exc}"
        model_load_wall_seconds = time.perf_counter() - load_started
        loaded = engine.loaded_info()
        if startup_error is None:
            if loaded.get("device") != "cuda":
                raise RuntimeError(f"Model loaded on {loaded.get('device')!r}, not CUDA.")
            if loaded.get("dtype") != "bfloat16":
                raise RuntimeError(f"Model dtype is {loaded.get('dtype')!r}, not BF16.")
            if str(Path(loaded.get("adapter") or "").resolve()) != str(
                    Path(config.ADAPTER).resolve()):
                raise RuntimeError("The expected LoRA adapter was not confirmed as loaded.")
            if not loaded.get("load_in_4bit"):
                raise RuntimeError("The validated 4-bit CUDA configuration was not loaded.")

        for position, case in enumerate(cases, start=1):
            doc_started = time.perf_counter()
            elapsed = time.perf_counter() - started
            print(f"[{position}/{len(cases)}] START {case['document_id']} "
                  f"(elapsed {elapsed:.1f}s)", flush=True)
            engine.reset_calls()
            payload = None
            meta: dict = {}
            failure = startup_error
            if failure is None:
                try:
                    rules = [{
                        "rule_id": rule.rule_id, "type": rule.pii_type,
                        "action": rule.action, "reason": rule.reason,
                    } for rule in lib.rules_for(case["receiver"])]
                    messages = prompts.analyze_messages(
                        case["text"], lib.receiver_label(case["receiver"]),
                        lib.receivers[case["receiver"]].purpose, rules)
                    payload, meta = engine.generate_json(
                        messages, purpose=f"evaluation {position}/{len(cases)} "
                        f"{case['document_id']}")
                    if payload is None:
                        attempts = int(meta.get("attempts") or 0)
                        failure = (f"No usable JSON answer after {attempts} attempt(s).")
                except Exception as exc:
                    failure = f"{type(exc).__name__}: {exc}"
                    torch.cuda.empty_cache()

            calls = engine.calls()
            total_in = sum(int(call.get("tokens_in") or 0) for call in calls)
            total_out = sum(int(call.get("tokens_out") or 0) for call in calls)
            inference_seconds = sum(float(call.get("seconds") or 0) for call in calls)
            elapsed_doc = time.perf_counter() - doc_started
            status = "failed" if failure else "completed"
            raw_record = {
                "document_id": case["document_id"],
                "status": status,
                "receiver": case["receiver"],
                "labels": case["items"],
                "prediction": payload,
                "generation_metadata": meta,
                "model_calls": calls,
                "error": failure,
            }
            raw_file.write(json.dumps(raw_record, ensure_ascii=False) + "\n")
            raw_file.flush()

            findings, shape_problems = metrics.parse_answer(payload)
            if shape_problems:
                error_rows.append({
                    "Document ID": case["document_id"],
                    "Error category": "Output format issue",
                    "PII value": None, "PII type": None, "Expected action": None,
                    "Predicted action": None, "Details": "; ".join(shape_problems),
                })
            if failure is None:
                matches = _match_findings(findings, case["items"])
                matched_finding_ids = set(matches.values())
                false_positives = [
                    (index, finding) for index, finding in enumerate(findings)
                    if metrics.is_structured(finding) and index not in matched_finding_ids
                ]
                correct_actions = 0
                checked_actions = 0
                for label_index, label in enumerate(case["items"]):
                    finding_index = matches.get(label_index)
                    expected = action_expected(label["type"], case["receiver"], lib)
                    finding = findings[finding_index] if finding_index is not None else None
                    predicted_action = (
                        str(finding.get("action", "")).upper() if finding else None)
                    action_correct = (
                        predicted_action == expected if finding is not None else None)
                    if action_correct is not None:
                        checked_actions += 1
                        correct_actions += int(action_correct)
                    pii_row = {
                        "Document ID": case["document_id"],
                        "Document status": status,
                        "Receiver": case["receiver"],
                        "PII value": label["text"],
                        "PII type": label["type"],
                        "Expected action": expected,
                        "Predicted value": finding.get("original_text") if finding else None,
                        "Predicted type": finding.get("type") if finding else None,
                        "Predicted action": predicted_action,
                        "Finding correct": finding is not None,
                        "Action correct": action_correct,
                        "Citation valid": (
                            metrics.citation_is_valid(finding, case["receiver"], lib)
                            if finding else None),
                        "Prediction reason": finding.get("reason") if finding else None,
                    }
                    pii_rows.append(pii_row)
                    if finding is None:
                        error_rows.append({
                            "Document ID": case["document_id"],
                            "Error category": "Missed PII",
                            "PII value": label["text"], "PII type": label["type"],
                            "Expected action": expected, "Predicted action": None,
                            "Details": "No structured finding of the correct PII type matched the label.",
                        })
                    elif not action_correct:
                        error_rows.append({
                            "Document ID": case["document_id"],
                            "Error category": "Incorrect decision",
                            "PII value": label["text"], "PII type": label["type"],
                            "Expected action": expected, "Predicted action": predicted_action,
                            "Details": str(finding.get("reason") or ""),
                        })
                    if label in must_not_be_visible(case["items"], case["receiver"], lib):
                        citation_valid = (
                            metrics.citation_is_valid(finding, case["receiver"], lib)
                            if finding else False)
                        if finding is None or predicted_action == "KEEP":
                            error_category = "Leak"
                        elif predicted_action == "ASK_HUMAN" or (
                                predicted_action in {"REPLACE", "REMOVE"} and not citation_valid):
                            error_category = "Deferred / invalid citation"
                        else:
                            error_category = None
                        if error_category:
                            error_rows.append({
                                "Document ID": case["document_id"],
                                "Error category": error_category,
                                "PII value": label["text"], "PII type": label["type"],
                                "Expected action": expected,
                                "Predicted action": predicted_action,
                                "Details": "The protected value would remain visible."
                                if finding is None or predicted_action == "KEEP"
                                else "The invalid citation causes the guard to defer this finding.",
                            })
                for _, finding in false_positives:
                    error_rows.append({
                        "Document ID": case["document_id"],
                        "Error category": "False positive",
                        "PII value": finding.get("original_text"),
                        "PII type": finding.get("type"),
                        "Expected action": None,
                        "Predicted action": finding.get("action"),
                        "Details": str(finding.get("reason") or
                                       "Structured finding did not match a labeled PII value."),
                    })
                score = metrics.score(
                    payload, case["items"], case["receiver"], library=lib,
                    document_id=case["document_id"])
                doc_scores.append(score)
                tp = len(matches)
                fp = len(false_positives)
                fn = len(case["items"]) - tp
                must_go_count = score.must_go
                leaked_count = len(score.leaked)
                deferred_count = len(score.deferred)
                doc_result = {
                    "Status": status, "Document ID": case["document_id"],
                    "Receiver": case["receiver"], "Labeled PII": len(case["items"]),
                    "Predicted findings": len(findings), "True positives": tp,
                    "False positives": fp, "Missed PII": fn,
                    "Precision": _ratio(tp, tp + fp), "Recall": _ratio(tp, tp + fn),
                    "F1": _ratio(2 * tp, 2 * tp + fp + fn),
                    "Action agreement": _ratio(correct_actions, checked_actions),
                    "Leak rate": _ratio(leaked_count, must_go_count),
                    "Deferred rate": _ratio(deferred_count, must_go_count),
                    "Inference seconds": round(inference_seconds, 3),
                    "Tokens in": total_in, "Tokens out": total_out, "Failure": None,
                }
            else:
                for label in case["items"]:
                    pii_rows.append({
                        "Document ID": case["document_id"],
                        "Document status": status, "Receiver": case["receiver"],
                        "PII value": label["text"], "PII type": label["type"],
                        "Expected action": action_expected(label["type"], case["receiver"], lib),
                        "Predicted value": None, "Predicted type": None,
                        "Predicted action": "NOT EVALUATED", "Finding correct": None,
                        "Action correct": None, "Citation valid": None,
                        "Prediction reason": None,
                    })
                error_rows.append({
                    "Document ID": case["document_id"], "Error category": "Document failure",
                    "PII value": None, "PII type": None, "Expected action": None,
                    "Predicted action": None, "Details": failure,
                })
                doc_result = {
                    "Status": status, "Document ID": case["document_id"],
                    "Receiver": case["receiver"], "Labeled PII": len(case["items"]),
                    "Predicted findings": None, "True positives": None,
                    "False positives": None, "Missed PII": None, "Precision": None,
                    "Recall": None, "F1": None, "Action agreement": None,
                    "Leak rate": None, "Deferred rate": None,
                    "Inference seconds": round(inference_seconds, 3),
                    "Tokens in": total_in, "Tokens out": total_out, "Failure": failure,
                }
            document_rows.append(doc_result)
            for call in calls:
                call_rows.append({
                    "Document ID": case["document_id"],
                    "Purpose": call.get("purpose"), "Device": call.get("device"),
                    "Seconds": call.get("seconds"), "Tokens in": call.get("tokens_in"),
                    "Tokens out": call.get("tokens_out"),
                    "Tokens/second": call.get("tokens_per_second"),
                    "Prompt chars": call.get("prompt_chars"),
                    "Output chars": call.get("output_chars"),
                    "JSON attempt": call.get("json_attempt"),
                    "JSON valid": call.get("json_valid"),
                    "JSON validation seconds": call.get("json_validation_seconds"),
                })
            elapsed = time.perf_counter() - started
            result_suffix = f"FAILED: {failure}" if failure else "completed"
            print(f"[{position}/{len(cases)}] {result_suffix}; "
                  f"document {elapsed_doc:.1f}s; total elapsed {elapsed:.1f}s",
                  flush=True)

    successful_documents = sum(row["Status"] == "completed" for row in document_rows)
    failed_documents = len(document_rows) - successful_documents
    evaluated_labels = sum(
        row["Labeled PII"] for row in document_rows if row["Status"] == "completed")
    tp = sum(int(row["True positives"] or 0) for row in document_rows
             if row["Status"] == "completed")
    fp = sum(int(row["False positives"] or 0) for row in document_rows
             if row["Status"] == "completed")
    fn = sum(int(row["Missed PII"] or 0) for row in document_rows
             if row["Status"] == "completed")
    all_findings = [
        finding
        for record in (json.loads(line) for line in json_path.read_text(encoding="utf-8").splitlines())
        if record["status"] == "completed"
        for finding in metrics.parse_answer(record["prediction"])[0]
    ]
    structured_count = sum(metrics.is_structured(finding) for finding in all_findings)
    citation_candidates = [
        finding for finding in all_findings
        if str(finding.get("action", "")).upper() in {"REPLACE", "REMOVE"}
    ]
    valid_citations = sum(
        metrics.citation_is_valid(finding, record["receiver"], lib)
        for record in (json.loads(line) for line in json_path.read_text(encoding="utf-8").splitlines())
        if record["status"] == "completed"
        for finding in metrics.parse_answer(record["prediction"])[0]
        if str(finding.get("action", "")).upper() in {"REPLACE", "REMOVE"}
    )
    protected_labels = [
        label for case in cases
        if any(row["Document ID"] == case["document_id"] and row["Status"] == "completed"
               for row in document_rows)
        for label in must_not_be_visible(case["items"], case["receiver"], lib)
    ]
    leaks = sum(row["Error category"] == "Leak" for row in error_rows)
    deferred = sum(row["Error category"] == "Deferred / invalid citation"
                   for row in error_rows)
    checked_actions = sum(
        row["Action correct"] is not None for row in pii_rows
        if row["Document status"] == "completed")
    correct_actions = sum(
        row["Action correct"] is True for row in pii_rows
        if row["Document status"] == "completed")
    runtime_seconds = time.perf_counter() - started
    gpu_props = torch.cuda.get_device_properties(0)
    loaded = engine.loaded_info()
    model_report = {
        "Base model": config.MODEL_NAME,
        "LoRA adapter": loaded.get("adapter") or config.ADAPTER,
        "Provider": engine.provider(),
        "Device": loaded.get("device"),
        "GPU": torch.cuda.get_device_name(0),
        "CUDA version": torch.version.cuda,
        "Dtype": loaded.get("dtype"),
        "Quantization": "NF4 4-bit" if loaded.get("load_in_4bit") else "none",
        "Attention": loaded.get("attention"),
        "GPU VRAM (GiB)": round(gpu_props.total_memory / 1024**3, 2),
        "Adapter loaded": bool(loaded.get("adapter")),
    }
    summary_metrics = [
        {"name": "Documents in dataset", "value": len(cases),
         "definition": "Number of source documents in the selected dataset."},
        {"name": "Successful documents", "value": successful_documents,
         "definition": "Documents with a completed model inference/evaluation."},
        {"name": "Failed documents", "value": failed_documents,
         "definition": "Documents whose inference raised an error; excluded from quality metric denominators."},
        {"name": "Labeled PII in dataset", "value": label_total,
         "definition": f"Ground-truth labeled entities across the {len(cases)} source documents."},
        {"name": "Labeled PII evaluated", "value": evaluated_labels,
         "definition": "Labels in successful documents; failed-document labels are not treated as misses."},
        {"name": "Labels unavailable due to failed documents",
         "value": label_total - evaluated_labels,
         "definition": "Ground-truth labels in failed documents, shown as not evaluated."},
        {"name": "True positives", "value": tp,
         "definition": "Structured findings matched one-to-one to a labeled value of the same PII type."},
        {"name": "False positives", "value": fp,
         "definition": "Structured predictions that did not match an unused label of the same PII type."},
        {"name": "Missed PII (false negatives)", "value": fn,
         "definition": "Labels in successful documents with no matching structured finding."},
        {"name": "PII Precision", "value": _ratio(tp, tp + fp),
         "numerator": tp, "denominator": tp + fp,
         "definition": "TP / (TP + FP): correctness of predicted detections.", "format": "percent"},
        {"name": "PII Recall", "value": _ratio(tp, tp + fn),
         "numerator": tp, "denominator": tp + fn,
         "definition": "TP / (TP + FN): fraction of evaluated labeled PII detected.", "format": "percent"},
        {"name": "PII F1", "value": _ratio(2 * tp, 2 * tp + fp + fn),
         "numerator": 2 * tp, "denominator": 2 * tp + fp + fn,
         "definition": "Harmonic mean of PII precision and recall.", "format": "percent"},
        {"name": "Action Agreement", "value": _ratio(correct_actions, checked_actions),
         "numerator": correct_actions, "denominator": checked_actions,
         "definition": "Among correctly matched labeled PII findings, fraction whose action equals the policy-derived expected action; separate from detection metrics.",
         "format": "percent"},
        {"name": "Structured output rate", "value": _ratio(structured_count, len(all_findings)),
         "numerator": structured_count, "denominator": len(all_findings),
         "definition": "Findings satisfying the evaluation schema / all parsed findings.",
         "format": "percent"},
        {"name": "Citation validity", "value": _ratio(valid_citations, len(citation_candidates)),
         "numerator": valid_citations, "denominator": len(citation_candidates),
         "definition": "Valid receiver-and-PII-type policy citations / REPLACE or REMOVE findings.",
         "format": "percent"},
        {"name": "Leak rate", "value": _ratio(leaks, len(protected_labels)),
         "numerator": leaks, "denominator": len(protected_labels),
         "definition": "Protected labels left visible because no finding matched or the model chose KEEP / labels required to be hidden in successful documents.",
         "format": "percent"},
        {"name": "Deferred rate", "value": _ratio(deferred, len(protected_labels)),
         "numerator": deferred, "denominator": len(protected_labels),
         "definition": "Protected labels deferred by ASK_HUMAN or invalid citation / protected labels in successful documents.",
         "format": "percent"},
        {"name": "Model load time (seconds)", "value": round(model_load_wall_seconds, 3),
         "definition": "Wall time for the one-time model and adapter load."},
        {"name": "Total inference time (seconds)",
         "value": round(sum(float(row.get("Seconds") or 0) for row in call_rows), 3),
         "definition": "Sum of measured model-call generation time; excludes non-generation pipeline work."},
        {"name": "Total runtime (seconds)", "value": round(runtime_seconds, 3),
         "definition": "Elapsed wall time from preflight through all documents and report metrics."},
    ]
    report = {
        "generated_at": generated_at, "json_path": json_path,
        "dataset_name": dataset_name, "test_dir": test_dir,
        "model": model_report, "summary_metrics": summary_metrics,
        "document_rows": document_rows, "pii_rows": pii_rows,
        "error_rows": error_rows, "call_rows": call_rows,
    }
    with xlsx_path.open("xb"):
        pass
    _write_workbook(xlsx_path, report)

    with json_path.open("r", encoding="utf-8") as handle:
        json_records = sum(1 for line in handle if line.strip())
    checked_book = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet_names = checked_book.sheetnames
    checked_book.close()
    expected_sheets = [
        "Summary", "Document Results", "PII-Level Results", "Errors", "Model Calls"]
    if json_records != len(cases) or sheet_names != expected_sheets:
        raise RuntimeError(
            f"Output verification failed: JSONL records={json_records}; sheets={sheet_names}.")
    print(f"Saved and verified {json_records} prediction records and workbook sheets "
          f"{', '.join(sheet_names)}.", flush=True)
    print(f"Excel report: {xlsx_path}", flush=True)
    print(f"Raw JSON predictions: {json_path}", flush=True)
    return xlsx_path, json_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--test-dir", type=Path, default=TEST_DIR)
    parser.add_argument("--output-dir", type=Path, default=RESULTS_DIR)
    parser.add_argument("--dataset-name")
    parser.add_argument("--expected-documents", type=int, default=EXPECTED_DOCUMENTS)
    parser.add_argument("--expected-labels", type=int, default=EXPECTED_LABELS)
    arguments = parser.parse_args()
    run(
        test_dir=arguments.test_dir,
        output_dir=arguments.output_dir,
        dataset_name=arguments.dataset_name,
        expected_documents=arguments.expected_documents,
        expected_labels=arguments.expected_labels,
    )
