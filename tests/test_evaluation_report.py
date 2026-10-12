from __future__ import annotations

import json
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from evaluation.full_report import _match_findings, _write_workbook
from evaluation.evaluate import test_cases as load_test_cases


ROOT = Path(__file__).resolve().parents[1]


def finding(value: str, pii_type: str = "PHONE", action: str = "REMOVE") -> dict:
    return {
        "original_text": value,
        "type": pii_type,
        "action": action,
        "replacement": None,
        "confidence": 0.95,
        "rule_id": "VENDOR-PII-003",
        "reason": "test",
    }


def test_label_matching_is_type_aware_and_one_to_one():
    labels = [{"text": "0501234567", "type": "PHONE"},
              {"text": "0501234567", "type": "PHONE"}]
    predictions = [finding("0501234567"), finding("0501234567"),
                   finding("0501234567", "NATIONAL_ID")]
    assert _match_findings(predictions, labels) == {0: 0, 1: 1}


def test_label_matching_accepts_a_finding_with_context():
    labels = [{"text": "0501234567", "type": "PHONE"}]
    predictions = [finding("Mobile number: 0501234567")]
    assert _match_findings(predictions, labels) == {0: 0}


def test_report_workbook_has_required_sheets_and_rows():
    report = {
        "generated_at": "2026-10-12T17:00:00+03:00",
        "json_path": Path("predictions.jsonl"),
        "dataset_name": "test_human_judge",
        "test_dir": Path("data/test_human_judge"),
        "model": {"Base model": "test"},
        "summary_metrics": [{
            "name": "PII Recall", "value": 1.0, "numerator": 1, "denominator": 1,
            "definition": "TP / (TP + FN).", "format": "percent",
        }],
        "document_rows": [{"Status": "completed", "Document ID": "doc-1"}],
        "pii_rows": [{"Document ID": "doc-1", "PII value": "masked-in-test"}],
        "error_rows": [],
        "call_rows": [],
    }
    with tempfile.TemporaryDirectory() as folder:
        target = Path(folder) / "report.xlsx"
        _write_workbook(target, report)
        workbook = load_workbook(target, read_only=True, data_only=True)
        try:
            assert workbook.sheetnames == [
                "Summary", "Document Results", "PII-Level Results", "Errors", "Model Calls"]
            assert workbook["Summary"]["A1"].value == \
                "SafeShare test_human_judge evaluation"
            assert Path(workbook["Summary"]["B3"].value) == Path("data/test_human_judge")
            assert workbook["PII-Level Results"].max_row == 2
            assert workbook["Document Results"].max_row == 2
        finally:
            workbook.close()


def test_human_judge_dataset_pairs_and_offsets_are_valid():
    dataset = ROOT / "data" / "test_human_judge"
    cases = load_test_cases(dataset)

    assert len(cases) == 20
    assert sum(len(case["items"]) for case in cases) == 87
    for gold_path in sorted(dataset.glob("*.gold.json")):
        gold = json.loads(gold_path.read_text(encoding="utf-8"))
        text_path = gold_path.with_name(gold_path.name.replace(".gold.json", ".txt"))
        assert text_path.is_file()
        text = text_path.read_text(encoding="utf-8")
        for span in gold["spans"]:
            assert text[span["start"]:span["end"]] == span["text"]
