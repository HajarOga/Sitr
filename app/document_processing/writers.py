"""Writing the released document.

The rule: change as little as possible. A DOCX is opened, the changed text is put
back into the same paragraph or table cell, and everything else - styles, headings,
table borders, numbers - is left exactly as it was. That is the difference between a
document someone can still use and a wall of plain text.

What is preserved: paragraphs, headings, tables and cells, page structure, styles.
What is not: formatting *inside* one paragraph (if a name was bold, the released value
takes the paragraph's first-run style), and PDF layout - PDF releases are written as
DOCX or TXT on purpose.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from .document import Document


def _set_paragraph_text(paragraph, new_text: str) -> None:
    """Replace a paragraph's text while keeping its first run's formatting."""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def write_docx(document: Document, out_path: str | Path, changed: dict[str, str]) -> Path:
    """Copy the original and patch in the released text.

    `changed` maps block id -> new text. Blocks that are not in it are untouched.
    """
    from docx import Document as DocxDocument            # noqa: PLC0415
    from docx.table import Table                          # noqa: PLC0415
    from docx.text.paragraph import Paragraph             # noqa: PLC0415

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    source = Path(document.path) if document.path else None
    original_bytes = document.meta.get("original_bytes")

    if original_bytes and (source is None or not source.exists()):
        # the upload was deleted after reading (that is the design), so write the
        # original bytes to a scratch file and patch that copy
        scratch = document.meta.get("scratch_path") or str(out_path.with_suffix(".src.docx"))
        Path(scratch).write_bytes(original_bytes)
        document.meta["scratch_path"] = scratch
        source = Path(scratch)

    if source is not None and source.suffix.lower() == ".docx" and source.exists():
        shutil.copyfile(source, out_path)
        docx = DocxDocument(str(out_path))
    else:
        # a PDF or TXT source: build a clean DOCX with the same text
        docx = DocxDocument()
        for block in document.blocks:
            text = changed.get(block.id, block.text)
            style = "Title" if block.kind == "heading" else None
            docx.add_paragraph(text, style=style) if style else docx.add_paragraph(text)
        docx.save(str(out_path))
        return out_path

    paragraphs = list(docx.paragraphs)
    tables = list(docx.tables)
    for block in document.blocks:
        if block.id not in changed:
            continue
        ref = block.ref
        try:
            if ref.get("kind") == "paragraph":
                _set_paragraph_text(paragraphs[ref["index"]], changed[block.id])
            elif ref.get("kind") == "cell":
                cell = tables[ref["table"]].rows[ref["row"]].cells[ref["col"]]
                _set_paragraph_text(cell.paragraphs[ref["para"]], changed[block.id])
        except (IndexError, KeyError):
            # a block we cannot locate is skipped rather than corrupting the file
            continue

    docx.save(str(out_path))
    return out_path


def write_txt(document: Document, out_path: str | Path, changed: dict[str, str]) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [changed.get(block.id, block.text) for block in document.blocks]
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path


def write_release(document: Document, out_dir: str | Path, stem: str,
                  changed: dict[str, str]) -> dict:
    """Write the release in the most faithful format available.

    DOCX in -> DOCX out. PDF in -> DOCX out (with a note), because PDF layout is not
    preserved by the prototype. TXT in -> TXT out.
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if document.kind == "docx":
        path = write_docx(document, out_dir / f"{stem}.safe.docx", changed)
        return {"format": "docx", "path": str(path), "layout_preserved": True}

    if document.kind == "pdf":
        path = write_docx(document, out_dir / f"{stem}.safe.docx", changed)
        return {"format": "docx", "path": str(path), "layout_preserved": False,
                "note": "converted from PDF: text and structure are kept, page layout is not"}

    path = write_txt(document, out_dir / f"{stem}.safe.txt", changed)
    return {"format": "txt", "path": str(path), "layout_preserved": True}


def safe_text(document: Document, changed: dict[str, str]) -> str:
    """The released text, used for display and for the leak-rate metric."""
    return "\n".join(changed.get(block.id, block.text) for block in document.blocks
                     if changed.get(block.id, block.text).strip())
